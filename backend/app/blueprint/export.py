"""Blueprint document export — editable DOCX (python-docx) and PDF (reportlab).

Renders the SAME assembled Blueprint into a professionally structured document
with headings and tables. No external services, no LLM — pure local rendering of
data already in state. Returns raw bytes for the API to stream.
"""
from __future__ import annotations

import io
from typing import List

from app.models.schemas import Blueprint


# --------------------------------------------------------------------------- #
# Section ordering shared by both renderers (title, kind, value-getter)
# --------------------------------------------------------------------------- #
def _sections(bp: Blueprint) -> list[tuple[str, str, object]]:
    idea = bp.selected_idea
    stack_rows = [[s.name, s.purpose] for s in bp.stack_details] or [
        [s, ""] for s in bp.recommended_stack
    ]
    team_rows = [
        [
            m.name,
            ", ".join(m.skills),
            "; ".join(m.responsibilities),
            ", ".join(m.technologies),
            m.time_allocation,
        ]
        for m in bp.team_allocation
    ]
    schedule_rows = [
        [s.task, s.technology, s.duration, s.dependency, s.expected_output]
        for s in bp.schedule
    ]
    return [
        ("Project Name", "text", bp.project_name or idea.title),
        ("Problem Statement", "text", bp.problem_statement),
        ("Hackathon Theme", "text", bp.hackathon_theme),
        ("Aim", "text", bp.aim),
        ("Objectives", "list", bp.objectives),
        ("Selected Idea", "text", idea.solution),
        ("Core Features", "list", bp.core_features),
        ("Target Users", "text", bp.target_users),
        ("Expected Impact", "text", bp.expected_impact),
        ("Technical Architecture", "list", bp.architecture),
        ("Technology Stack", "table", (["Technology", "Purpose"], stack_rows)),
        (
            "Team & Responsibilities",
            "table",
            (["Name", "Skills", "Responsibilities", "Technologies", "Time"], team_rows),
        ),
        (
            "Development Timeline",
            "table",
            (["Task", "Technology", "Duration", "Dependency", "Expected Output"], schedule_rows),
        ),
        ("Implementation Procedure", "list", bp.implementation_procedure),
        ("Dependencies", "list", bp.dependencies),
        ("Technical Risks", "list", bp.technical_risks),
        ("Risk Mitigation", "list", bp.risk_mitigation),
        ("Testing Plan", "list", bp.testing_plan),
        ("Demo Flow (2-minute)", "list", bp.demo_flow),
        ("Value Proposition", "text", bp.value_proposition),
        ("Pitch Structure", "list", bp.pitch_structure),
        ("Future Scope", "list", bp.future_scope),
        ("MVP Features", "list", bp.mvp_features),
        ("Nice-to-Have Features", "list", bp.nice_to_have_features),
    ]


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
def blueprint_to_docx(bp: Blueprint) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("Hackathon Project Blueprint", level=0)
    if bp.alignment_report:
        doc.add_paragraph(
            f"Overall alignment: {bp.alignment_report.overall}/100 "
            f"({bp.alignment_report.label})"
        )

    for title, kind, value in _sections(bp):
        doc.add_heading(title, level=1)
        if kind == "text":
            doc.add_paragraph(str(value or "—"))
        elif kind == "list":
            items = value or []
            if not items:
                doc.add_paragraph("—")
            for it in items:
                doc.add_paragraph(str(it), style="List Bullet")
        elif kind == "table":
            headers, rows = value
            if not rows:
                doc.add_paragraph("—")
                continue
            table = doc.add_table(rows=1, cols=len(headers))
            # Not every base template ships this named style; fall back gracefully
            # instead of raising KeyError (which would surface as an opaque 500).
            try:
                table.style = "Light Grid Accent 1"
            except (KeyError, ValueError):
                try:
                    table.style = "Table Grid"
                except (KeyError, ValueError):
                    pass
            for i, h in enumerate(headers):
                cell = table.rows[0].cells[i]
                cell.text = h
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.bold = True
                        run.font.size = Pt(9)
            for row in rows:
                cells = table.add_row().cells
                for i, val in enumerate(row):
                    cells[i].text = str(val or "")
                    for p in cells[i].paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def blueprint_to_pdf(bp: Blueprint) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        ListFlowable,
        ListItem,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
    )
    story: list = [Paragraph("Hackathon Project Blueprint", styles["Title"])]
    if bp.alignment_report:
        story.append(
            Paragraph(
                f"Overall alignment: {bp.alignment_report.overall}/100 "
                f"({bp.alignment_report.label})",
                body,
            )
        )
    story.append(Spacer(1, 6))

    def esc(s: object) -> str:
        return (
            str(s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        )

    for title, kind, value in _sections(bp):
        story.append(Paragraph(esc(title), styles["Heading2"]))
        if kind == "text":
            story.append(Paragraph(esc(value) or "—", body))
        elif kind == "list":
            items = value or []
            if items:
                story.append(
                    ListFlowable(
                        [ListItem(Paragraph(esc(it), body)) for it in items],
                        bulletType="bullet",
                    )
                )
            else:
                story.append(Paragraph("—", body))
        elif kind == "table":
            headers, rows = value
            if rows:
                data = [[Paragraph(esc(h), body) for h in headers]] + [
                    [Paragraph(esc(c), body) for c in row] for row in rows
                ]
                table = Table(data, repeatRows=1, hAlign="LEFT")
                table.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ]
                    )
                )
                story.append(table)
            else:
                story.append(Paragraph("—", body))
        story.append(Spacer(1, 4))

    doc.build(story)
    return buf.getvalue()
